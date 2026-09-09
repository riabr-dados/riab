"""Gera downloads integrais e manifesto público a partir dos Parquets locais.

CSV em partições de até 500 mil linhas; XLSX para recursos de até 100 mil linhas.
O manifesto descreve somente arquivos efetivamente gerados e reabertos.
"""
from pathlib import Path
from decimal import Decimal
from datetime import datetime, date
import argparse
import csv
import gzip
import hashlib
import json
import os
import time

import openpyxl
from openpyxl.cell import WriteOnlyCell
import pyarrow.parquet as pq
import yaml

ROOT = Path(__file__).resolve().parents[1]
CSV_ROWS = 500_000
XLSX_ROWS = 100_000


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def xlsx_cell(sheet, value):
    if value is None:
        return None
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, Decimal):
        # Excel limita números a 15 dígitos; valores mais precisos ficam como texto.
        value = str(value) if len(value.as_tuple().digits) > 15 else float(value)
    if isinstance(value, int) and len(str(abs(value))) > 15:
        value = str(value)
    if isinstance(value, str):
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = 's'
        cell.number_format = '@'
        return cell
    return value


def generate(path, destination, *, csv_rows=CSV_ROWS, xlsx_rows=XLSX_ROWS, gzip_rows=CSV_ROWS):
    source = pq.ParquetFile(path)
    rows = source.metadata.num_rows
    columns = source.schema_arrow.names
    if len(columns) != len(set(columns)):
        raise ValueError(f'{path}: colunas duplicadas')
    destination.mkdir(parents=True, exist_ok=True)
    descriptor = {'rows': rows, 'columns': len(columns), 'parquet_sha256': digest(path), 'files': []}
    book = openpyxl.Workbook(write_only=True) if rows <= xlsx_rows and len(columns) <= 16384 else None
    if book:
        sheet = book.create_sheet('Dados')
        sheet.freeze_panes = 'A2'
        sheet.append([xlsx_cell(sheet, c) for c in columns])
    part, n_part, total = 0, 0, 0
    stream = None

    def close_csv():
        nonlocal stream
        stream.close()
        # Conferir CSV sem depender da representação intermediária do Parquet.
        opener = gzip.open if csv_path.suffix == '.gz' else open
        with opener(csv_path, mode='rt', encoding='utf-8-sig', newline='') as verify:
            reader = csv.reader(verify)
            if next(reader) != columns or sum(1 for _ in reader) != n_part:
                raise ValueError(f'CSV incompleto: {csv_path}')
        item = {'format': 'csv', 'path': csv_path.name, 'rows': n_part,
                'bytes': csv_path.stat().st_size, 'sha256': digest(csv_path)}
        if csv_path.suffix == '.gz':
            item['compression'] = 'gzip'
        descriptor['files'].append(item)
        stream = None

    try:
        batches = source.iter_batches(batch_size=10000)
        for batch in batches:
            for record in batch.to_pylist():
                if stream is None:
                    part += 1
                    suffix = f'.part-{part:04d}' if rows > csv_rows else ''
                    extension = '.csv.gz' if rows > gzip_rows else '.csv'
                    csv_path = destination / f'{path.stem}{suffix}{extension}'
                    stream = (gzip.open(csv_path, 'wt', encoding='utf-8-sig', newline='')
                              if extension.endswith('.gz') else
                              csv_path.open('w', encoding='utf-8-sig', newline=''))
                    writer = csv.writer(stream)
                    writer.writerow(columns)
                    n_part = 0
                values = [record[c] for c in columns]
                writer.writerow([json.dumps(v, ensure_ascii=False, default=str) if isinstance(v, (list, dict)) else v for v in values])
                if book:
                    sheet.append([xlsx_cell(sheet, v) for v in values])
                total += 1
                n_part += 1
                if n_part == csv_rows:
                    close_csv()
        if stream is not None:
            close_csv()
        if total != rows:
            raise ValueError(f'{path}: leitura incompleta')
        if rows == 0:
            csv_path = destination / f'{path.stem}.csv'
            stream = csv_path.open('w', encoding='utf-8-sig', newline='')
            csv.writer(stream).writerow(columns)
            n_part = 0
            close_csv()
        if book:
            from openpyxl.utils import get_column_letter
            sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{rows + 1}'
            dictionary = book.create_sheet('Dicionário')
            dictionary.append(['Campo', 'Tipo Parquet'])
            for field in source.schema_arrow:
                dictionary.append([field.name, str(field.type)])
            dictionary.append(['Precisão Excel', 'Identificadores e números com mais de 15 dígitos significativos são texto para preservar o valor.'])
            output = destination / f'{path.stem}.xlsx'
            book.save(output)
            check = openpyxl.load_workbook(output, read_only=True, data_only=True)
            count = sum(1 for _ in check['Dados'].values) - 1
            check.close()
            if count != rows:
                raise ValueError(f'XLSX incompleto: {output}')
            descriptor['files'].append({'format': 'xlsx', 'path': output.name, 'rows': rows,
                                       'bytes': output.stat().st_size, 'sha256': digest(output)})
        else:
            descriptor['xlsx_note'] = 'Use os CSVs completos por partição; XLSX não gerado devido ao tamanho da tabela.'
    finally:
        if stream is not None:
            stream.close()
    return descriptor


def catalog_tables(catalog_path):
    """Lista somente recursos publicados, aceitando entradas simples e detalhadas."""
    datasets = yaml.safe_load(catalog_path.read_text(encoding='utf-8'))['datasets']
    tables = []
    for dataset in datasets:
        if dataset.get('hidden'):
            continue
        for resource in dataset.get('cleaned', {}).get('tables', []):
            tables.append(resource if isinstance(resource, str) else resource['name'])
    if len(tables) != len(set(tables)):
        raise ValueError('O catálogo referencia o mesmo recurso mais de uma vez')
    return tables


def descriptor_is_current(path, destination, descriptor):
    if not descriptor or descriptor.get('parquet_sha256') != digest(path):
        return False
    if descriptor.get('rows', 0) > CSV_ROWS and any(
        item.get('format') == 'csv' and item.get('compression') != 'gzip'
        for item in descriptor.get('files', [])
    ):
        return False
    for item in descriptor.get('files', []):
        output = destination / item.get('path', '')
        if not output.is_file() or output.stat().st_size != item.get('bytes') or digest(output) != item.get('sha256'):
            return False
    return bool(descriptor.get('files'))


def save_manifest(path, manifest, *, attempts=8, delay=0.25):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.json.tmp')
    temp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for attempt in range(attempts):
        try:
            os.replace(temp, path)
            return
        except PermissionError:
            if attempt + 1 == attempts:
                raise
            time.sleep(delay * (attempt + 1))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--tables', nargs='*')
    parser.add_argument('--input-dir', type=Path, default=ROOT / 'pipelines/output/cleaned')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'pipelines/output/downloads')
    parser.add_argument('--manifest', type=Path, default=ROOT / 'catalog/downloads.json')
    parser.add_argument('--catalog', type=Path, default=ROOT / 'catalog/datasets.yaml')
    args = parser.parse_args()
    existing = json.loads(args.manifest.read_text(encoding='utf-8')) if args.manifest.exists() else {'version': 1, 'resources': {}}
    tables = args.tables if args.tables else catalog_tables(args.catalog)
    files = [args.input_dir / f'{table}.parquet' for table in tables]
    if not files:
        raise ValueError('Nenhum recurso para exportar')
    missing = [str(path) for path in files if not path.exists()]
    if missing:
        raise ValueError(f'Recursos do catálogo sem Parquet local: {missing}')
    manifest = {'version': 1, 'resources': {}}
    for path in files:
        descriptor = existing.get('resources', {}).get(path.stem)
        if descriptor_is_current(path, args.output_dir, descriptor):
            manifest['resources'][path.stem] = descriptor
            print(f'{path.stem}: download atual reutilizado')
        else:
            manifest['resources'][path.stem] = generate(path, args.output_dir)
            print(f'{path.stem}: {manifest["resources"][path.stem]["rows"]} linhas exportadas')
        # Checkpoint atômico: uma interrupção não perde os recursos já verificados.
        save_manifest(args.manifest, manifest)


if __name__ == '__main__':
    main()
